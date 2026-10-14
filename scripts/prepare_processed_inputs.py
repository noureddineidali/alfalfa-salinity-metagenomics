#!/usr/bin/env python3
"""Prepare and validate project metadata and supplementary-data tables."""

from __future__ import annotations

import csv
from collections import Counter
from decimal import Decimal, getcontext
from pathlib import Path

from openpyxl import load_workbook


getcontext().prec = 30

ROOT = Path(__file__).resolve().parents[1]
REFERENCE_DIR = ROOT / "references" / "lv_2026"
METADATA_DIR = ROOT / "data" / "metadata"
PROCESSED_DIR = ROOT / "data" / "processed"
DOCS_DIR = ROOT / "docs"

SAMPLE_IDS = [
    "CK-1", "CK-2", "CK-3",
    "LS-1", "LS-2", "LS-3",
    "MS-1", "MS-2", "MS-3",
]

TREATMENTS = {
    "CK": "Control",
    "LS": "Low salinity",
    "MS": "Moderate salinity",
}

ACCESSIONS = {
    "CK-1": ("SAMN54495940", "SRS27706072", "SRX31738113", "SRR36739085"),
    "CK-2": ("SAMN54495941", "SRS27706074", "SRX31738114", "SRR36739084"),
    "CK-3": ("SAMN54495942", "SRS27706073", "SRX31738115", "SRR36739083"),
    "LS-1": ("SAMN54495943", "SRS27706075", "SRX31738116", "SRR36739082"),
    "LS-2": ("SAMN54495944", "SRS27706076", "SRX31738117", "SRR36739081"),
    "LS-3": ("SAMN54495945", "SRS27706077", "SRX31738118", "SRR36739080"),
    "MS-1": ("SAMN54495946", "SRS27706078", "SRX31738119", "SRR36739079"),
    "MS-2": ("SAMN54495947", "SRS27706080", "SRX31738120", "SRR36739078"),
    "MS-3": ("SAMN54495948", "SRS27706079", "SRX31738121", "SRR36739077"),
}

ENV_SOURCE_HEADERS = [
    "",
    "EC",
    "pH",
    "TN",
    "Nitrate nitrogen",
    "Ammonium nitrogen",
    "Plant height",
    "Fresh weight",
    "Dry weight",
]

ENV_OUTPUT_HEADERS = [
    "sample_id",
    "treatment",
    "replicate",
    "ec_us_cm",
    "ph",
    "total_n_g_kg",
    "nitrate_n_mg_kg",
    "ammonium_n_mg_kg",
    "plant_height_cm",
    "fresh_weight_g",
    "dry_weight_g",
]

DOMINANT_TAXA = [
    "SCGC-AG-212-J23",
    "Gp6-AA40",
    "Sediminibacterium",
    "PSRF01",
    "12-FULL-67-14b",
    "Bradyrhizobium",
    "UBA12294",
    "JACDDX01",
    "WHTF01",
    "Nocardioides",
]

MAPPING_TOLERANCE_PERCENTAGE_POINTS = Decimal("0.00000005")


def treatment_and_replicate(sample_id: str) -> tuple[str, int]:
    prefix, replicate_text = sample_id.split("-", maxsplit=1)
    if prefix not in TREATMENTS or replicate_text not in {"1", "2", "3"}:
        raise ValueError(f"Invalid sample identifier: {sample_id}")
    return TREATMENTS[prefix], int(replicate_text)


def write_tsv(path: Path, header: list[str], rows: list[list[object]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter="\t", lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)


def prepare_manifest() -> list[list[object]]:
    if list(ACCESSIONS) != SAMPLE_IDS:
        raise ValueError("Manifest accessions do not contain the expected ordered samples")

    rows = []
    for sample_id in SAMPLE_IDS:
        treatment, replicate = treatment_and_replicate(sample_id)
        rows.append([sample_id, treatment, replicate, *ACCESSIONS[sample_id]])

    if len(rows) != 9 or len({row[0] for row in rows}) != 9:
        raise ValueError("The manifest must contain exactly nine unique samples")

    write_tsv(
        METADATA_DIR / "sample_manifest.tsv",
        [
            "sample_id",
            "treatment",
            "replicate",
            "biosample_accession",
            "sra_sample_accession",
            "experiment_accession",
            "run_accession",
        ],
        rows,
    )
    return rows


def prepare_environmental_metadata() -> list[list[object]]:
    source = REFERENCE_DIR / "data sheet 1.csv"
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        source_rows = list(reader)

    if header != ENV_SOURCE_HEADERS:
        raise ValueError(f"Unexpected data sheet 1.csv header: {header}")
    if [row[0] for row in source_rows] != SAMPLE_IDS:
        raise ValueError("data sheet 1.csv samples are missing or out of expected order")
    if any(len(row) != len(ENV_SOURCE_HEADERS) for row in source_rows):
        raise ValueError("Unexpected row width in data sheet 1.csv")

    output_rows = []
    for row in source_rows:
        sample_id = row[0]
        treatment, replicate = treatment_and_replicate(sample_id)
        output_rows.append([sample_id, treatment, replicate, *row[1:]])

    write_tsv(
        PROCESSED_DIR / "environmental_plant_metadata.tsv",
        ENV_OUTPUT_HEADERS,
        output_rows,
    )
    return output_rows


def read_dominant_genus_csv() -> tuple[list[str], dict[str, dict[str, Decimal]]]:
    source = REFERENCE_DIR / "data sheet 2.csv"
    with source.open("r", encoding="utf-8-sig", newline="") as handle:
        reader = csv.reader(handle)
        header = next(reader)
        rows = list(reader)

    if header[0] != "" or header[1:-1] != DOMINANT_TAXA or header[-1] != "Others":
        raise ValueError("Unexpected data sheet 2.csv structure")
    if [row[0] for row in rows] != SAMPLE_IDS:
        raise ValueError("data sheet 2.csv does not contain the expected labelled samples")

    values = {
        row[0]: {taxon: Decimal(value) for taxon, value in zip(header[1:], row[1:])}
        for row in rows
    }
    return header[1:], values


def read_genus_workbook() -> tuple[list[str], list[list[str]]]:
    source = REFERENCE_DIR / "table 4.xlsx"
    workbook = load_workbook(source, read_only=True, data_only=True)
    if workbook.sheetnames != ["Sheet1"]:
        raise ValueError(f"Unexpected table 4.xlsx worksheets: {workbook.sheetnames}")
    sheet = workbook["Sheet1"]
    iterator = sheet.iter_rows(values_only=True)
    header_values = next(iterator)
    headers = [str(value) if value is not None else "" for value in header_values]
    if headers != ["genus", "1", "2", "3", "4", "5", "6", "7", "8", "9"]:
        raise ValueError(f"Unexpected table 4.xlsx header: {headers}")

    rows: list[list[str]] = []
    for excel_row in iterator:
        if len(excel_row) != 10:
            raise ValueError("Unexpected row width in table 4.xlsx")
        rows.append(["" if value is None else str(value) for value in excel_row])
    workbook.close()
    return headers, rows


def verify_column_mapping(
    workbook_rows: list[list[str]],
    labelled_values: dict[str, dict[str, Decimal]],
) -> tuple[dict[int, str], list[Decimal], Decimal]:
    totals = [
        sum((Decimal(row[column]) for row in workbook_rows if row[column] != ""), Decimal(0))
        for column in range(1, 10)
    ]
    if any(total == 0 for total in totals):
        raise ValueError("A table 4.xlsx sample column has a zero total")

    by_taxon = {row[0]: row for row in workbook_rows}
    missing_reference_taxa = [taxon for taxon in DOMINANT_TAXA if taxon not in by_taxon]
    if missing_reference_taxa:
        raise ValueError(f"Dominant taxa absent from table 4.xlsx: {missing_reference_taxa}")

    mapping: dict[int, str] = {}
    matched_errors: list[Decimal] = []
    for column in range(1, 10):
        errors_by_sample = {}
        for sample_id in SAMPLE_IDS:
            errors = []
            for taxon in DOMINANT_TAXA:
                workbook_percent = Decimal(100) * Decimal(by_taxon[taxon][column]) / totals[column - 1]
                errors.append(abs(workbook_percent - labelled_values[sample_id][taxon]))
            errors_by_sample[sample_id] = max(errors)

        candidates = [
            sample_id
            for sample_id, error in errors_by_sample.items()
            if error <= MAPPING_TOLERANCE_PERCENTAGE_POINTS
        ]
        if len(candidates) != 1:
            raise ValueError(
                f"Workbook column {column} does not uniquely match a labelled sample: {candidates}"
            )
        sample_id = candidates[0]
        mapping[column] = sample_id
        matched_errors.append(errors_by_sample[sample_id])

    if len(set(mapping.values())) != 9:
        raise ValueError("The inferred workbook-column mapping is not one-to-one")
    expected_mapping = {index: sample_id for index, sample_id in enumerate(SAMPLE_IDS, start=1)}
    if mapping != expected_mapping:
        raise ValueError(f"Verified mapping differs from expected sample order: {mapping}")

    return mapping, totals, max(matched_errors)


def prepare_genus_abundance() -> dict[str, object]:
    _, labelled_values = read_dominant_genus_csv()
    _, workbook_rows = read_genus_workbook()

    labels = [row[0] for row in workbook_rows]
    label_counts = Counter(labels)
    missing_labels = sum(label.strip() == "" for label in labels)
    duplicate_groups = {label: count for label, count in label_counts.items() if count > 1}
    missing_values = sum(value == "" for row in workbook_rows for value in row[1:])

    if len(workbook_rows) != 15805:
        raise ValueError(f"Unexpected genus row count: {len(workbook_rows)}")
    if missing_labels or duplicate_groups:
        raise ValueError("Missing or duplicated genus labels prevent safe conversion")

    mapping, totals, maximum_error = verify_column_mapping(workbook_rows, labelled_values)
    ordered_columns = [next(column for column, sample in mapping.items() if sample == sample_id) for sample_id in SAMPLE_IDS]
    output_rows = [
        [row[0], *[row[column] for column in ordered_columns]]
        for row in workbook_rows
    ]
    write_tsv(
        PROCESSED_DIR / "genus_abundance.tsv",
        ["genus", *SAMPLE_IDS],
        output_rows,
    )

    return {
        "rows": len(workbook_rows),
        "columns": 10,
        "mapping": mapping,
        "totals": totals,
        "maximum_error": maximum_error,
        "dominant_comparisons": len(DOMINANT_TAXA) * len(SAMPLE_IDS),
        "missing_labels": missing_labels,
        "duplicate_groups": len(duplicate_groups),
        "missing_values": missing_values,
    }


def prepare_pathway_abundance() -> dict[str, object]:
    source = REFERENCE_DIR / "table 7.xls"
    source_bytes = source.read_bytes()
    try:
        source_text = source_bytes.decode("ascii")
    except UnicodeDecodeError as error:
        raise ValueError("table 7.xls is not ASCII text as expected") from error

    if source_bytes.startswith((b"\xd0\xcf\x11\xe0", b"PK\x03\x04")):
        raise ValueError("table 7.xls unexpectedly contains an Excel workbook")
    if "\t" not in source_text:
        raise ValueError("table 7.xls is not tab-delimited")

    with source.open("r", encoding="ascii", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        source_header = next(reader)
        source_rows = list(reader)

    expected_source_header = ["", *SAMPLE_IDS]
    if source_header != expected_source_header:
        raise ValueError(f"Unexpected table 7.xls header: {source_header}")
    if len(source_rows) != 30:
        raise ValueError(f"Unexpected pathway row count: {len(source_rows)}")
    if any(len(row) != len(expected_source_header) for row in source_rows):
        raise ValueError("Unexpected row width in table 7.xls")

    pathway_names = [row[0] for row in source_rows]
    pathway_counts = Counter(pathway_names)
    missing_pathway_names = sum(name.strip() == "" for name in pathway_names)
    duplicate_pathways = {name: count for name, count in pathway_counts.items() if count > 1}
    missing_values = sum(value == "" for row in source_rows for value in row[1:])

    numeric_values: list[Decimal] = []
    for row_number, row in enumerate(source_rows, start=2):
        for column_number, value in enumerate(row[1:], start=2):
            if value == "":
                continue
            try:
                numeric_values.append(Decimal(value))
            except Exception as error:
                raise ValueError(
                    f"Non-numeric pathway value at row {row_number}, column {column_number}: {value}"
                ) from error

    negative_values = sum(value < 0 for value in numeric_values)
    if missing_pathway_names or duplicate_pathways:
        raise ValueError("Missing or duplicated pathway labels prevent safe conversion")
    if missing_values:
        raise ValueError("Missing pathway abundance values prevent lossless conversion")

    output_path = PROCESSED_DIR / "pathway_abundance.tsv"
    output_header = ["pathway", *SAMPLE_IDS]
    write_tsv(output_path, output_header, source_rows)

    with output_path.open("r", encoding="utf-8", newline="") as handle:
        reader = csv.reader(handle, delimiter="\t")
        written_header = next(reader)
        written_rows = list(reader)

    if written_header != output_header:
        raise ValueError("Processed pathway header does not match the required header")
    if [row[0] for row in written_rows] != pathway_names:
        raise ValueError("Processed pathway names or row order differ from the source")
    if len(written_rows) != len(source_rows):
        raise ValueError("Processed pathway row count differs from the source")
    numeric_values_match_exactly = all(
        written_row[1:] == source_row[1:]
        for written_row, source_row in zip(written_rows, source_rows)
    )
    if not numeric_values_match_exactly:
        raise ValueError("Processed pathway values differ textually from the source")

    return {
        "source_format": "ASCII tab-delimited text (not a binary Excel workbook)",
        "rows": len(source_rows),
        "columns": len(output_header),
        "column_names": output_header,
        "missing_pathway_names": missing_pathway_names,
        "duplicate_pathways": len(duplicate_pathways),
        "missing_values": missing_values,
        "negative_values": negative_values,
        "numeric_values": len(numeric_values),
        "numeric_values_match_exactly": numeric_values_match_exactly,
    }


def write_validation_report(
    genus_stats: dict[str, object],
    pathway_stats: dict[str, object],
) -> None:
    totals = genus_stats["totals"]
    mapping = genus_stats["mapping"]
    report = f"""# Processed data validation

The four processed inputs were generated by `scripts/prepare_processed_inputs.py` from the accession information in `docs/dataset_reconnaissance.md`, `references/lv_2026/data sheet 1.csv`, `references/lv_2026/data sheet 2.csv`, `references/lv_2026/table 4.xlsx`, and `references/lv_2026/table 7.xls`. Source supplementary files were read without modification.

## Sample structure

- `sample_manifest.tsv` contains exactly 9 unique samples: CK-1–CK-3, LS-1–LS-3, and MS-1–MS-3.
- Treatment assignments are Control (CK), Low salinity (LS), and Moderate salinity (MS), with replicates 1–3 in each treatment.
- `environmental_plant_metadata.tsv` contains the same 9 samples and 8 original measured variables. It has 9 data rows and 11 columns after adding sample, treatment, and replicate fields.

## Genus matrix

- Dimensions: {genus_stats['rows']:,} genus rows × 9 sample columns ({genus_stats['columns']} columns including the genus label).
- The numbered workbook columns were verified against the ten explicitly labelled dominant genera in `data sheet 2.csv`: {', '.join(f'{column}={mapping[column]}' for column in range(1, 10))}.
- All {genus_stats['dominant_comparisons']} genus–sample percentage comparisons agreed within {MAPPING_TOLERANCE_PERCENTAGE_POINTS} percentage points. The largest absolute difference was {genus_stats['maximum_error']:.12f} percentage points.
- Missing genus labels: {genus_stats['missing_labels']}.
- Duplicated genus labels: {genus_stats['duplicate_groups']}.
- Missing abundance cells: {genus_stats['missing_values']}.

### Preserved abundance-column totals

| Sample | Column total |
|---|---:|
{chr(10).join(f'| {sample_id} | {total:.6f} |' for sample_id, total in zip(SAMPLE_IDS, totals))}

No abundance values were normalized, transformed, filtered, aggregated, or statistically analysed. Percentages were calculated only in memory to verify the sample-column mapping and were not written to the genus matrix.

## Pathway matrix

- Source format: {pathway_stats['source_format']}.
- Dimensions: {pathway_stats['rows']} pathway rows × 9 sample columns ({pathway_stats['columns']} columns including the pathway label).
- Column names: {', '.join(f'`{name}`' for name in pathway_stats['column_names'])}.
- The 30 pathway names and their original row order were preserved exactly; only the blank first-column header was renamed `pathway`.
- All {pathway_stats['numeric_values']} numerical cells match the source text exactly.
- Missing pathway names: {pathway_stats['missing_pathway_names']}.
- Duplicated pathway names: {pathway_stats['duplicate_pathways']}.
- Missing abundance cells: {pathway_stats['missing_values']}.
- Negative abundance values: {pathway_stats['negative_values']}.

Validation passed. Values were not transformed, converted to percentages, normalized, scaled, rounded, filtered, reordered, renamed, or imputed. The source does not define the abundance unit, so the values remain labelled as pathway abundance without a more specific unit.
"""
    DOCS_DIR.mkdir(parents=True, exist_ok=True)
    (DOCS_DIR / "processed_data_validation.md").write_text(report, encoding="utf-8")


def main() -> None:
    prepare_manifest()
    prepare_environmental_metadata()
    genus_stats = prepare_genus_abundance()
    pathway_stats = prepare_pathway_abundance()
    write_validation_report(genus_stats, pathway_stats)


if __name__ == "__main__":
    main()
